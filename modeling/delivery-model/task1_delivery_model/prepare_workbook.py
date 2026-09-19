from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'.python_packages'))
import pandas as pd
import openpyxl
w=openpyxl.load_workbook(ROOT/'output/task1_evidence_and_results.xlsx')
(ROOT/'results/workbook_sheets.json').write_text(json.dumps(w.sheetnames),encoding='utf-8')
mapping={'Parameters':'calibration/calibrated_parameters.csv','Baseline 24h':'baseline/baseline_24h_endpoints.csv',
 'Benchmark Metrics':'calibration/all_benchmark_metrics.csv','Uncertainty Summary':'uncertainty/uncertainty_summary.csv',
 'Sensitivity':'uncertainty/sensitivity_spearman.csv','Numerical Tests':'numerical_tests.csv',
 'Grid Convergence':'grid_convergence.csv','Task2 Interface':'task2_interface.csv',
 'Uncertainty Inputs':'uncertainty/uncertainty_parameter_sets_256.csv','Uncertainty Raw':'uncertainty/uncertainty_256x4.csv'}
blocks=[]
def block(sheet,path,row=4,columns=None,title=None,note=None):
 f=pd.read_csv(path)
 if columns is not None:f=f[columns]
 matrix=[list(f.columns)]+json.loads(f.to_json(orient='values'))
 blocks.append(dict(sheet=sheet,row=row,matrix=matrix,title=title,note=note))
for name,file in mapping.items():
 f=pd.read_csv(ROOT/'results'/file)
 columns=list(f.columns)
 if name in ['Baseline 24h','Uncertainty Raw']:
  columns=[v for v in next(w[name].iter_rows(min_row=4,max_row=4,values_only=True)) if v]
  if 'target_retention_ug_cm2' not in columns: columns+=['target_retention_ug_cm2']
 block(name,ROOT/'results'/file,columns=columns)
block('Scope Gate',ROOT/'data/processed/scope_gate.csv')
block('Evidence Matrix',ROOT/'data/processed/evidence_matrix.csv')
block('Scope Gate',ROOT/'results/acceptance_audit.csv',17)
block('Uncertainty Summary',ROOT/'results/uncertainty/pairwise_differences.csv',13)
block('Target Exposure',ROOT/'results/secondary_endpoints.csv',title='Layer exposure and transport readouts',note='VE and dermis concentration AUC are reported separately (ug h/cm3).')
block('Human Transfer',ROOT/'results/decision/human_transfer_metrics.csv',title='Human-assisted transfer scenarios',note='Human data constrain conditional parameters; assisted fitting is not independent validation.')
block('Human Transfer',ROOT/'results/decision/human_identifiability.csv',10)
block('Boundary Conditions',ROOT/'results/decision/boundary_consistency_metrics.csv',title='Boundary conditions and effective shape calibration',note='Standardized geometry and donor assumptions; fluorescence shape does not identify absolute uptake.')
block('Boundary Conditions',ROOT/'results/decision/boundary_domain_scenarios.csv',15)
block('Decision Conditions',ROOT/'results/decision/break_even_surface_P.csv',title='Equal-retention conditions',note='tFNA surface-P roots with D fixed; comparator and geometry held at baseline. Empty root means no crossing in the scanned range.')
(ROOT/'results/workbook_data.json').write_text(json.dumps(blocks,ensure_ascii=False),encoding='utf-8')
