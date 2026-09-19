"""Cross-check scientific tables and delivered report references."""
from pathlib import Path
import sys,re,json,py_compile
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'.python_packages'))
import numpy as np
import pandas as pd
import openpyxl
from pypdf import PdfReader

rows=[]
def check(name,passed):
    rows.append(dict(check=name,passed=bool(passed)))
    if not passed:raise AssertionError(name)
for p in (ROOT/'src').glob('*.py'):py_compile.compile(str(p),doraise=True)
for suffix in ['完整','简要','图解']:
    p=ROOT/'output'/f'任务一_递送建模横向比较{suffix}报告.md'
    text=p.read_text(encoding='utf-8')
    check(suffix+'_finished_prose',not re.search('修订|修改说明|已修改|不再|旧版|新版|前后对比',text))
    for target in re.findall(r'\]\(([^)]+)\)',text):
        if not target.startswith(('http','#')):check(suffix+'_link_'+target,(p.parent/target).resolve().exists())
    pdftext='\n'.join(page.extract_text() or '' for page in PdfReader(p.with_suffix('.pdf')).pages)
    check(suffix+'_pdf_primary_value','0.5864' in pdftext)
    check(suffix+'_pdf_has_text',len(pdftext)>1000)
for file in ['baseline/baseline_24h_endpoints.csv','uncertainty/uncertainty_256x4.csv']:
    f=pd.read_csv(ROOT/'results'/file)
    check(file+'_target_identity',np.allclose(f.target_retention_ug_cm2,f.VE_mass_ug_cm2+f.dermis_mass_ug_cm2,atol=1e-10))
f=pd.read_csv(ROOT/'results/decision/decision_surface.csv')
check('decision_grid_441',len(f)==441)
for file in ['numerical_tests.csv','acceptance_audit.csv','decision/decision_checks.csv']:
    check(file,pd.read_csv(ROOT/'results'/file).passed.all())
w=openpyxl.load_workbook(ROOT/'output/task1_evidence_and_results.xlsx',data_only=True)
baseline=pd.read_csv(ROOT/'results/baseline/baseline_24h_endpoints.csv')
for i,row in baseline.iterrows():
    check('xlsx_primary_'+row.system,abs(w['Overview'].cell(i+5,2).value-row.target_retention_ug_cm2)<1e-8)
check('xlsx_numeric_checks',w['Overview']['B18'].value==18 and w['Overview']['B20'].value==10)
check('xlsx_sheet_names_unique',len(w.sheetnames)==len(set(w.sheetnames)))
for sheet in w:
    check('xlsx_errors_'+sheet.title,not any(c.data_type=='e' for row in sheet for c in row))
(ROOT/'results/document_checks/delivery_checks.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Delivery checks passed: {len(rows)}')
